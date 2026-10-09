"""Generate the table fixtures in tests/fixtures/inputs/ (they are also checked in).

Run from the repo root to regenerate:  .venv\\Scripts\\python.exe -m tests.inputs.fixturegen

The base is the user's sample trackid.net export, fjaak-techno-dj-set-section-may-2026.csv,
copied unchanged from the v1 folder. Every variant is derived from it, so the same 17 tracks
must come out of each one. CSV/TXT outputs are byte-deterministic (tests compare them);
the .xlsx files are compared by content (zip entries carry timestamps).
"""
from __future__ import annotations

import codecs
import csv
import datetime as dt
import io
import unicodedata
from pathlib import Path

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "inputs"
SAMPLE = "fjaak-techno-dj-set-section-may-2026.csv"

# (file name, delimiter, encoding, line end) -- the brief's 2 delimiters x 3 encodings, plus the
# tab-separated UTF-16 "Texto Unicode" that Excel also writes. "," + utf-8 + LF is the sample.
VARIANTS = (
    ("fjaak-coma-utf8-bom.csv", ",", "utf-8-sig", "\r\n"),
    ("fjaak-coma-cp1252.csv", ",", "cp1252", "\r\n"),
    ("fjaak-puntoycoma-cp1252.csv", ";", "cp1252", "\r\n"),       # Excel en español, "CSV"
    ("fjaak-puntoycoma-utf8-bom.csv", ";", "utf-8-sig", "\r\n"),  # Excel, "CSV UTF-8"
    ("fjaak-puntoycoma-utf8.csv", ";", "utf-8", "\n"),
    ("fjaak-tab-utf16.txt", "\t", "utf-16", "\r\n"),               # Excel, "Texto Unicode"
)


def sample_rows(dst: Path = FIXTURES) -> list[list[str]]:
    text = (dst / SAMPLE).read_bytes().decode("utf-8")
    return [r for r in csv.reader(io.StringIO(text)) if r]


def csv_bytes(rows: list[list], delimiter: str, encoding: str, line_end: str) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, delimiter=delimiter, lineterminator=line_end).writerows(rows)
    text = buf.getvalue()
    if encoding == "utf-16":
        return codecs.BOM_UTF16_LE + text.encode("utf-16-le")
    return text.encode(encoding)


def _hms(s: str) -> dt.time:
    h, m, sec = (int(x) for x in s.split(":"))
    return dt.time(h, m, sec)


def write_xlsx(path: Path, rows: list[list], *, time_cols: tuple[int, ...] = (), sheet: str = "Tracklist",
               blank_first_sheet: bool = False) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    if blank_first_sheet:
        ws.title = "Vacia"
        ws = wb.create_sheet(sheet)
    else:
        ws.title = sheet
    for r in rows:
        ws.append(r)
    for col in time_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col + 1, max_col=col + 1):
            cell.number_format = "hh:mm:ss"
    fixed = dt.datetime(2026, 10, 9, 0, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(path)


def small_tables() -> dict[str, bytes]:
    nfd = lambda s: unicodedata.normalize("NFD", s)  # noqa: E731
    rows_120 = [["Pos", "Artist", "Title"]] + [[str(i), f"Artista {i:03d}", f"Titulo {i:03d}"]
                                                for i in range(1, 121)]
    return {
        "sin-pos.csv": csv_bytes([["Artist", "Title"], ["Phara", "Grey Ribbon"], ["Klint", "Gunslingers"],
                                  ["Exos", "Glaður Og Reifur"]], ",", "utf-8", "\n"),
        "MAYUSCULAS.CSV": csv_bytes([["Artist", "Title"], ["Phara", "Grey Ribbon"], ["Exos", "Glaður Og Reifur"]],
                                    ";", "cp1252", "\r\n"),
        # sandbox case D: a row added by hand without Pos
        "pos-incompleto.csv": csv_bytes([["Pos", "Artist", "Title"], ["1", "Phara", "Grey Ribbon"],
                                         ["2", "Klint", "Gunslingers"], ["", "Exos", "Glaður Og Reifur"]],
                                        ",", "utf-8", "\n"),
        # sandbox case B: an ID row and an empty title
        "titulo-vacio.csv": csv_bytes([["Pos", "Artist", "Title"], ["1", "Phara", "Grey Ribbon"],
                                       ["2", "ID", "ID"], ["3", "Ferdinand", ""], ["4", "Klint", "Gunslingers"]],
                                      ",", "utf-8", "\n"),
        "nfd.csv": csv_bytes([["Artist", "Title"], [nfd("Óscar Mulero"), nfd("Cuerpo Sólido")],
                              [nfd("Rødhåd"), nfd("Wise Saying")]], ",", "utf-8", "\n"),
        "apostrofo.csv": csv_bytes([["Artist", "Title"], ["Ferdinger", "I\u2019m Still Here"],
                                    ["Stef Mendesidis", "The Secret Agent\u2019s Blunder"]], ",", "utf-8", "\n"),
        "artistas-numericos.csv": csv_bytes([["Artist", "Title"], ["808 State", "Pacific 202"],
                                             ["2 Unlimited", "No Limit"], ["Andrea Belluzzi", "75"]],
                                            ",", "utf-8", "\n"),
        "mas-de-99.csv": csv_bytes(rows_120, ",", "utf-8", "\n"),
        "una-columna.csv": csv_bytes([["Track"], ["Truncate - Dust 1 (Kai Van Dongen Remix) [Truncate]"],
                                      ["Funk Tribu ft. LUCY DYE - All Of It"], ["Gunslingers"]], ",", "utf-8", "\n"),
        "sin-cabecera.csv": csv_bytes([["Phara", "Grey Ribbon"], ["Klint", "Gunslingers"]], ",", "utf-8", "\n"),
        # DJ-software export: "Time" is the track length, not the start in the set
        "export-dj.csv": csv_bytes([["#", "Track Title", "Artist", "Time"], ["1", "Grey Ribbon", "Phara", "05:09"],
                                    ["2", "Gunslingers", "Klint", "06:41"]], ",", "utf-8", "\n"),
        "tracklist.txt": ("Tracklist\r\n00:00 Phara - Grey Ribbon\r\n05:10 Klint - Gunslingers\r\n"
                          "11:00 ID - ID\r\n").encode("utf-8"),
    }


def generate(dst: Path = FIXTURES) -> list[Path]:
    """Write every derived fixture into dst (which must already hold the sample CSV)."""
    rows = sample_rows(dst)
    out = []
    for name, delimiter, encoding, line_end in VARIANTS:
        p = dst / name
        p.write_bytes(csv_bytes(rows, delimiter, encoding, line_end))
        out.append(p)
    for name, data in small_tables().items():
        p = dst / name
        p.write_bytes(data)
        out.append(p)
    xrows = [rows[0]] + [[int(r[0]), _hms(r[1]), _hms(r[2]), *r[3:]] for r in rows[1:]]
    write_xlsx(dst / "fjaak.xlsx", xrows, time_cols=(1, 2))
    out.append(dst / "fjaak.xlsx")
    write_xlsx(dst / "mi-lista.xlsx", [
        ["Nº", "Artista", "Título", "Versión", "Sello"],
        [1, "Truncate", "Dust 1", "Kai Van Dongen Remix", "Truncate"],
        [2.0, 1991, 75, None, None],
        [None, "Exos", "Glaður Og Reifur", None, "Mutual Rytm"],
    ], sheet="Mi set", blank_first_sheet=True)
    out.append(dst / "mi-lista.xlsx")
    return out


if __name__ == "__main__":
    for p in generate():
        print(p.name, p.stat().st_size)
